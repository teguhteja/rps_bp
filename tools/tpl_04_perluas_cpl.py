#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Perbesar kapasitas CPL pada template: baris daftar CPL dan kolom matriks
"Korelasi CPL terhadap Sub-CPMK".

Dibutuhkan karena Tugas Akhir (TSI7854) membebani 17 CPL, sementara template
hasil `tpl_01`-`tpl_03` hanya menyediakan 16.

Skrip idempoten: kalau kapasitasnya sudah mencukupi, tidak ada yang diubah.

Pemakaian:
    python tools/tpl_04_perluas_cpl.py [RPS_MK_TSI0000.docx] [kapasitas]
"""
import copy
import re
import sys

from docx import Document
from docx.oxml.ns import qn

BAKU = 20
LEBAR_LABEL, LEBAR_TEPI = 1700, 1000
LEBAR_ISI = 13782                       # lebar sel penampung (gridSpan 12)


def teks_sel(tc):
    return ''.join(t.text or '' for t in tc.iter(qn('w:t')))


def indeks_terbesar(tbl, nama):
    angka = re.findall(nama + r'\[(\d+)\]',
                       ''.join(t.text or '' for t in tbl.iter(qn('w:t'))))
    return max(int(a) for a in angka) if angka else -1


def retarget(el, indeks):
    for t in el.iter(qn('w:t')):
        if t.text and '[' in t.text:
            t.text = re.sub(r'\[\s*\d+\s*\]', '[%d]' % indeks, t.text)


def perluas_baris(tbl, kapasitas):
    """
    Gandakan baris daftar CPL terakhir sampai jumlah slotnya cukup.

    Baris daftar dikenali dari placeholder `cpl_prodi[n].label`. Placeholder
    `cpl_prodi[n].nomor` sengaja diabaikan karena itu milik header matriks
    Korelasi CPL — kalau ikut terhitung, baris matriksnyalah yang tergandakan.
    """
    baris_daftar = {}
    for tr in tbl.findall(qn('w:tr')):
        teks = ''.join(t.text or '' for t in tr.iter(qn('w:t')))
        for a in re.findall(r'cpl_prodi\[(\d+)\]\.label', teks):
            baris_daftar[int(a)] = tr
    if not baris_daftar:
        return 0
    terakhir = max(baris_daftar)
    if terakhir + 1 >= kapasitas:
        return 0
    sumber = baris_daftar[terakhir]
    anchor, n = sumber, 0
    for i in range(terakhir + 1, kapasitas):
        baru = copy.deepcopy(sumber)
        anchor.addnext(baru)
        retarget(baru, i)
        anchor = baru
        n += 1
    return n


def cari_matriks(doc):
    """Tabel bersarang matriks Korelasi CPL (header memuat 'Bobot penilaian')."""
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for nested in cell.tables:
                    kepala = ' '.join(c.text for c in nested.rows[0].cells)
                    if 'Bobot penilaian' in kepala:
                        return nested
    return None


def perluas_kolom(nested, kapasitas):
    """Sisipkan kolom CPL sebelum kolom Bobot dan Jumlah Minggu."""
    tbl = nested._tbl
    terakhir = indeks_terbesar(tbl, 'cpl_prodi')
    jumlah = terakhir + 1
    if jumlah >= kapasitas:
        return 0
    tambah = kapasitas - jumlah

    # Sel matriks memuat DUA angka — `korelasi_cpl[baris].p<kolom>` — sehingga
    # tidak boleh dipakaikan retarget() yang menimpa semua angka dalam kurung.
    # Teks tiap sel salinan ditulis ulang secara eksplisit sesuai posisinya.
    for r_i, tr in enumerate(tbl.findall(qn('w:tr'))):
        tcs = tr.findall(qn('w:tc'))
        sumber = tcs[jumlah]              # kolom CPL terakhir
        anchor = sumber
        for k in range(tambah):
            kolom = jumlah + k            # indeks CPL berbasis 0
            baru = copy.deepcopy(sumber)
            anchor.addnext(baru)
            teks = ('{cpl_prodi[%d].nomor}' % kolom if r_i == 0
                    else '{korelasi_cpl[%d].p%d}' % (r_i - 1, kolom + 1))
            simpul = list(baru.iter(qn('w:t')))
            if simpul:
                simpul[0].text = teks
                for s in simpul[1:]:
                    s.text = ''
            anchor = baru

    grid = tbl.find(qn('w:tblGrid'))
    kolom = grid.findall(qn('w:gridCol'))
    for k in range(tambah):
        grid.insert(list(grid).index(kolom[jumlah]) + 1, copy.deepcopy(kolom[jumlah]))

    # bagi ulang lebar: label + N kolom CPL + Bobot + Jumlah Minggu
    sisa = (LEBAR_ISI - LEBAR_LABEL - 2 * LEBAR_TEPI) // kapasitas
    lebar = [LEBAR_LABEL] + [sisa] * kapasitas + [LEBAR_TEPI, LEBAR_TEPI]
    for c, w in zip(grid.findall(qn('w:gridCol')), lebar):
        c.set(qn('w:w'), str(w))
    for tr in tbl.findall(qn('w:tr')):
        for tc, w in zip(tr.findall(qn('w:tc')), lebar):
            tcPr = tc.find(qn('w:tcPr'))
            tcW = tcPr.find(qn('w:tcW')) if tcPr is not None else None
            if tcW is not None:
                tcW.set(qn('w:w'), str(w))
                tcW.set(qn('w:type'), 'dxa')
    return tambah


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else 'template/RPS_MK_TSI0000.docx'
    kapasitas = int(sys.argv[2]) if len(sys.argv) > 2 else BAKU
    doc = Document(path)

    n_baris = perluas_baris(doc.tables[4]._tbl, kapasitas)
    matriks = cari_matriks(doc)
    n_kolom = perluas_kolom(matriks, kapasitas) if matriks is not None else 0
    if matriks is None:
        print('PERINGATAN: matriks Korelasi CPL tidak ditemukan.')

    doc.save(path)
    print('kapasitas CPL -> %d (baris +%d, kolom matriks +%d)'
          % (kapasitas, n_baris, n_kolom))


if __name__ == '__main__':
    main()
