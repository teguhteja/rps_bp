#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Tampilkan placeholder yang masih tersisa di dokumen DOCX hasil generate.

Placeholder yang tersisa berarti nama field di JSON tidak cocok dengan
template. Untuk pemeriksaan JSON sebelum digenerate, pakai
`tools/validasi_rps.py`.

Pemakaian:
    python check_placeholders.py outputs/rps/sem1.1/TSI1107_Agama.docx
    python check_placeholders.py "outputs/rps/sem*/*.docx"

Status keluar 1 bila ada placeholder tersisa.
"""
import glob
import os
import re
import sys

from docx import Document

POLA = re.compile(r'\{[^{}]+\}')


def cari(doc_path):
    doc = Document(doc_path)
    ditemukan = set()

    def dari_paragraf(p):
        ditemukan.update(POLA.findall(p.text))

    def dari_tabel(t):
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    dari_paragraf(p)
                for nested in cell.tables:
                    dari_tabel(nested)

    for p in doc.paragraphs:
        dari_paragraf(p)
    for t in doc.tables:
        dari_tabel(t)
    for s in doc.sections:
        for bagian in (s.header, s.footer):
            for p in bagian.paragraphs:
                dari_paragraf(p)
            for t in bagian.tables:
                dari_tabel(t)
    return sorted(ditemukan)


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)

    berkas = []
    for arg in sys.argv[1:]:
        berkas.extend(sorted(glob.glob(arg)) or ([arg] if os.path.exists(arg) else []))
    if not berkas:
        raise SystemExit('Tidak ada berkas cocok dengan: %s' % ' '.join(sys.argv[1:]))

    bermasalah = 0
    for f in berkas:
        sisa = cari(f)
        if sisa:
            bermasalah += 1
            print('%s -> %d placeholder tersisa' % (f, len(sisa)))
            for p in sisa:
                print('    ', p)
        else:
            print('%s -> bersih' % f)
    print('\n%d berkas diperiksa, %d bermasalah' % (len(berkas), bermasalah))
    return 1 if bermasalah else 0


if __name__ == '__main__':
    sys.exit(main())
