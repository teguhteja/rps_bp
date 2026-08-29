# -*- coding: utf-8 -*-
"""
Selaraskan RPS_MK_TSI0000.docx dengan format acuan
'materi-rps/KFR 4560 Cosmeceutical.docx'.

Perubahan:
  1. Baris CPL      : 3 kolom (kode | tipe | deskripsi) -> 2 kolom (label | deskripsi)
                      label berbentuk 'CPL 1 (S-01)'.
  2. Baris CPMK     : 3 kolom -> 2 kolom; rujukan CPL pindah ke dalam deskripsi.
  3. Baris Sub-CPMK : 5 kolom -> 2 kolom; taksonomi, CPMK, dan minggu pindah
                      ke dalam deskripsi.
  4. Tambah band 'Penilaian dan Korelasinya dengan CPMK' berisi matriks
     jenis penilaian x bobot x kontribusi tiap CPMK.
  5. Header kolom (4) tabel mingguan: 'Kriteria & Bentuk' -> 'Kriteria & Teknik'.
"""
import copy, re, sys
from docx import Document
from docx.table import _Cell
from docx.oxml.ns import qn
from docx.oxml import parse_xml, OxmlElement
from docx.shared import Pt, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH

MAX_CPMK, MAX_PENILAIAN = 10, 10
HOLDER_WIDTH = 13782

doc = Document(sys.argv[1] if len(sys.argv) > 1 else 'template/RPS_MK_TSI0000.docx')
tbl = doc.tables[4]


def span_of(tc):
    tcPr = tc.find(qn('w:tcPr'))
    gs = tcPr.find(qn('w:gridSpan')) if tcPr is not None else None
    return int(gs.get(qn('w:val'))) if gs is not None else 1


def set_span(tc, nilai):
    tcPr = tc.find(qn('w:tcPr'))
    gs = tcPr.find(qn('w:gridSpan'))
    if gs is None:
        gs = OxmlElement('w:gridSpan')
        tcPr.insert(0, gs)
    gs.set(qn('w:val'), str(nilai))


def tulis(tc, teks):
    """Isi sel dengan satu paragraf berisi teks (format run pertama dipakai)."""
    paras = tc.findall(qn('w:p'))
    from docx.text.paragraph import Paragraph
    p = Paragraph(paras[0], None)
    if p.runs:
        p.runs[0].text = teks
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run(teks)
    for extra in paras[1:]:
        tc.remove(extra)


def ciutkan(prefix, kolom_isi, label_baru=None):
    """Sisakan sel label kiri, sel kode, dan satu sel deskripsi lebar."""
    diubah = 0
    for row in tbl.rows:
        teks = ''.join(row._tr.itertext())
        if prefix not in teks:
            continue
        tcs = row._tr.findall(qn('w:tc'))
        if len(tcs) <= 3:
            continue
        kode_tc, isi_tc = tcs[1], tcs[2]
        lebar_total = sum(span_of(tc) for tc in tcs[2:])
        for tc in tcs[3:]:
            row._tr.remove(tc)
        set_span(isi_tc, lebar_total)
        idx = re.search(r'\[(\d+)\]', ''.join(kode_tc.itertext()))
        if idx and label_baru:
            tulis(kode_tc, label_baru % int(idx.group(1)))
        if idx:
            tulis(isi_tc, '{%s[%d].deskripsi}' % (kolom_isi, int(idx.group(1))))
        diubah += 1
    return diubah


n_cpl = ciutkan('{cpl_prodi[', 'cpl_prodi', '{cpl_prodi[%d].label}')
n_cpmk = ciutkan('{cpmk [', 'cpmk')
n_sub = ciutkan('{sub_cpmk[', 'sub_cpmk')
print('baris diciutkan -> CPL %d, CPMK %d, Sub-CPMK %d' % (n_cpl, n_cpmk, n_sub))

# --- band Penilaian dan Korelasinya, disisipkan sebelum 'Deskrisi Singkat MK'
anchor = None
for row in tbl.rows:
    if 'Deskrisi' in ''.join(row._tr.itertext()):
        anchor = row._tr
        break
if anchor is None:
    raise SystemExit("baris 'Deskrisi Singkat MK' tidak ditemukan")

band = copy.deepcopy(anchor)
for h in (band.find(qn('w:trPr')) or []):
    if h.tag == qn('w:trHeight'):
        band.find(qn('w:trPr')).remove(h)
tcs = band.findall(qn('w:tc'))
tulis(tcs[0], 'Penilaian dan Korelasinya dengan CPMK')
for p_el in tcs[1].findall(qn('w:p')):
    for t in p_el.iter(qn('w:t')):
        t.text = ''
anchor.addprevious(band)

cell = _Cell(tcs[1], tbl)
judul = cell.paragraphs[0]
run = judul.add_run('Korelasi CPMK terhadap penilaian')
run.bold = True

nested = cell.add_table(rows=MAX_PENILAIAN + 1, cols=MAX_CPMK + 2)
nested.autofit = False
nested._tbl.tblPr.append(parse_xml(
    '<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    + ''.join('<w:%s w:val="single" w:sz="4" w:space="0" w:color="000000"/>' % e
              for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'))
    + '</w:tblBorders>'))

lebar_jenis, lebar_bobot = 4200, 900
sisa = (HOLDER_WIDTH - lebar_jenis - lebar_bobot) // MAX_CPMK
for r_i, row in enumerate(nested.rows):
    for c_i, sel in enumerate(row.cells):
        sel.width = Twips(lebar_jenis if c_i == 0 else lebar_bobot if c_i == 1 else sisa)
        if r_i == 0:
            teks = 'Penilaian' if c_i == 0 else 'Bobot' if c_i == 1 else '{cpmk[%d].kode}' % (c_i - 2)
        elif c_i == 0:
            teks = '{penilaian[%d].jenis}' % (r_i - 1)
        elif c_i == 1:
            teks = '{penilaian[%d].bobot}' % (r_i - 1)
        else:
            teks = '{penilaian[%d].c%d}' % (r_i - 1, c_i - 1)
        p = sel.paragraphs[0]
        if c_i > 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(teks)
        r.font.size = Pt(8)
        if r_i == 0:
            r.bold = True
print('band penilaian ditambahkan: %d x %d' % (MAX_PENILAIAN + 1, MAX_CPMK + 2))

# --- header kolom (4) tabel mingguan
n = 0
for row in doc.tables[5].rows[:3]:
    for t in row._tr.iter(qn('w:t')):
        if t.text and 'Bentuk' in t.text and 'Kriteria' in ''.join(row._tr.itertext()):
            baru = t.text.replace('Kriteria & Bentuk', 'Kriteria & Teknik')
            if baru != t.text:
                t.text = baru
                n += 1
print('header (4) diganti jadi "Kriteria & Teknik":', n)

doc.save(sys.argv[1] if len(sys.argv) > 1 else 'template/RPS_MK_TSI0000.docx')
print('template disimpan')
